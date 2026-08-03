#!/usr/bin/env python3
"""Replace the placeholder TOC in a DOCX with a static, page-numbered TOC.

Page numbers are read from the already-rendered PDF (same basename) so the
TOC in the re-rendered PDF shows real page numbers. Assumes the TOC occupies
a single page in both passes (true when entry count is small), so body page
numbers do not shift between passes.

Usage: python3 inject_toc.py <docx> <pdf> [maxlevel]
"""
import os, re, sys, tempfile, zipfile
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from pypdf import PdfReader

def norm(s):
    return re.sub(r'[^a-z0-9]', '', s.lower())


def update_document_statistics(docx_path, pages, text_parts):
    """Populate the extended properties that python-docx leaves at 1/0."""
    full_text = "\n".join(part for part in text_parts if part)
    values = {
        "Pages": pages,
        "Words": len(re.findall(r"\b\w[\w'-]*\b", full_text, flags=re.UNICODE)),
        "Characters": len(re.sub(r"\s", "", full_text)),
        "CharactersWithSpaces": len(full_text),
        "Lines": max(1, len(full_text.splitlines())),
        "Paragraphs": sum(1 for part in text_parts if part.strip()),
    }
    with zipfile.ZipFile(docx_path, "r") as source:
        members = [(item, source.read(item.filename)) for item in source.infolist()]
    fd, temp_path = tempfile.mkstemp(suffix=".docx", dir=os.path.dirname(os.path.abspath(docx_path)))
    os.close(fd)
    try:
        with zipfile.ZipFile(temp_path, "w") as target:
            for item, data in members:
                if item.filename == "docProps/app.xml":
                    xml = data.decode("utf-8")
                    for tag, value in values.items():
                        xml = re.sub(
                            fr"(<{tag}>).*?(</{tag}>)",
                            fr"\g<1>{value}\g<2>",
                            xml,
                            count=1,
                        )
                    data = xml.encode("utf-8")
                target.writestr(item, data)
        os.replace(temp_path, docx_path)
    finally:
        if os.path.exists(temp_path):
            os.unlink(temp_path)

def main():
    docx_path, pdf_path = sys.argv[1], sys.argv[2]
    maxlevel = int(sys.argv[3]) if len(sys.argv) > 3 else 2

    doc = Document(docx_path)

    # 1) collect heading paragraphs (in order) up to maxlevel
    heads = []
    for p in doc.paragraphs:
        sn = p.style.name
        if sn.startswith("Heading"):
            try:
                lvl = int(sn.split()[-1])
            except ValueError:
                continue
            if lvl <= maxlevel and p.text.strip():
                heads.append((lvl, p.text.strip()))

    # 2) page text from the rendered PDF
    reader = PdfReader(pdf_path)
    pages = [norm(pg.extract_text() or "") for pg in reader.pages]

    # 3) Locate the body after the rendered TOC. LibreOffice may populate the
    #    placeholder field in pass 1, so searching from page 1 would match the
    #    heading names inside the TOC itself rather than their body pages.
    toc_page = next((i for i, page in enumerate(pages) if norm("TABLE OF CONTENTS") in page), None)
    body_start = 0
    if toc_page is not None and heads:
        first_key = norm(heads[0][1])[:24]
        body_start = next(
            (i for i in range(toc_page + 1, len(pages)) if first_key and first_key in pages[i]),
            toc_page + 1,
        )

    # 4) map each heading -> first body page (>= cursor) whose text contains it
    entries = []
    cursor = body_start
    for lvl, text in heads:
        key = norm(text)[:24]
        page = None
        for i in range(cursor, len(pages)):
            if key and key in pages[i]:
                page = i + 1  # 1-based, matches footer "Page N"
                cursor = i
                break
        if page is None:  # fallback: search from the body start, never the TOC
            for i in range(body_start, len(pages)):
                if key and key in pages[i]:
                    page = i + 1
                    break
        entries.append((lvl, text, page or 1))

    # 4b) Compensate only when the static TOC page count differs from the TOC
    #     already rendered in pass 1. In normal LibreOffice builds both counts
    #     match, so body page numbers remain unchanged.
    LINES_PER_PAGE = 44
    toc_lines = len(entries) + 2            # "Contents" heading + spacing
    toc_pages = max(1, -(-toc_lines // LINES_PER_PAGE))
    existing_toc_pages = max(1, body_start - toc_page) if toc_page is not None else 1
    offset = toc_pages - existing_toc_pages
    if offset:
        entries = [(lvl, text, pg + offset) for lvl, text, pg in entries]

    # 5) find the placeholder TOC paragraph (contains a TOC field or the hint text)
    placeholder = None
    for p in doc.paragraphs:
        xml = p._p.xml
        if ("Update Field" in p.text) or ("TOC " in xml and 'instrText' in xml) or ("fldSimple" in xml and "TOC" in xml):
            placeholder = p
            break
    if placeholder is None:
        text_parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            text_parts.extend(cell.text for row in table.rows for cell in row.cells)
        update_document_statistics(docx_path, len(pages), text_parts)
        print("  [inject_toc] static TOC already present; refreshed document statistics in", docx_path)
        return

    # 6) build static TOC paragraphs, insert before placeholder, then delete it
    anchor = placeholder._p
    GREY = RGBColor(0x33, 0x33, 0x33)
    for lvl, text, page in entries:
        new_p = anchor.makeelement(qn('w:p'), {})
        anchor.addprevious(new_p)
        from docx.text.paragraph import Paragraph
        para = Paragraph(new_p, placeholder._parent)
        pf = para.paragraph_format
        pf.tab_stops.add_tab_stop(Inches(6.3), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        if lvl >= 2:
            pf.left_indent = Inches(0.3)
        pf.space_after = Pt(3)
        r = para.add_run(text + "\t" + str(page))
        r.font.size = Pt(11 if lvl == 1 else 10.5)
        r.font.name = "Arial"
        r.bold = (lvl == 1)
        r.font.color.rgb = GREY
    anchor.getparent().remove(anchor)

    text_parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        text_parts.extend(cell.text for row in table.rows for cell in row.cells)
    doc.save(docx_path)
    update_document_statistics(docx_path, len(pages) + offset, text_parts)
    print(f"  [inject_toc] {docx_path}: wrote {len(entries)} TOC entries "
          f"(pages {entries[0][2]}..{entries[-1][2]})")

if __name__ == "__main__":
    main()
